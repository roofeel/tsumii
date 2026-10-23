import csv
import io

from openpyxl import load_workbook
from PIL import Image

from app.engines.catalog import load_sample_catalog
from app.engines.motifs import render_motif
from app.exports.common import SheetSource
from app.exports.image import pattern_png, pattern_svg
from app.exports.pdf import pattern_pdf
from app.exports.tables import shopping_csv, shopping_xlsx

CAT = load_sample_catalog()
CELLS = render_motif("leaf", ["s01", "s14", "s15", "s16", "s12"], 20, 24)
SRC = SheetSource(name="Matcha Garden 吊坠", item_type="pendant", width=20, height=24, cells=CELLS, style_tags=["清新", "植物"], difficulty="beginner",
                  feasibility_note="含立体花朵结构，MVP 只能转换平面部分")


def test_csv_has_all_colors_and_bom():
    data = shopping_csv(SRC, CAT)
    assert data.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    codes = [r[0] for r in rows if r and r[0].startswith("SAMPLE-")]
    assert len(codes) == len(set(CELLS))
    assert sum(int(r[3]) for r in rows if r and r[0].startswith("SAMPLE-")) == len(CELLS)


def test_xlsx_formulas_and_sheets():
    wb = load_workbook(io.BytesIO(shopping_xlsx(SRC, CAT)))
    assert wb.sheetnames == ["珠子材料", "其他配件"]
    ws = wb["珠子材料"]
    assert ws["D4"].value == "使用数量"
    assert ws["E5"].value.startswith("=CEILING") and ws["F5"].value == "=D5+E5"


def test_png_and_svg():
    png = pattern_png(CELLS, 20, 24, CAT, {c: i + 1 for i, c in enumerate(sorted(set(CELLS)))})
    im = Image.open(io.BytesIO(png))
    assert im.size == (34 + 20 * 26 + 16, 34 + 24 * 26 + 16)
    svg = pattern_svg(CELLS, 20, 24, CAT)
    assert svg.startswith("<svg") and svg.count("<rect") == 20 * 24 + 1


def test_pdf_builds_multiple_pages():
    pdf = pattern_pdf(SRC, CAT)
    assert pdf.startswith(b"%PDF") and len(pdf) > 10_000
    pages = pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages")
    assert pages >= 3
