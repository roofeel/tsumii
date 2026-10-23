"""采购清单 CSV / XLSX 导出"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engines.catalog import Catalog
from ..engines.materials import calculate_materials, estimate_accessories
from .common import STATUS_ZH, SheetSource

BEAD_HEADER = ["色号", "颜色名称", "系列规格", "使用数量", "建议备量", "需要购买", "每包颗数", "包装数量", "已购买", "缺货", "数据状态"]
PART_HEADER = ["配件", "数量", "说明", "状态"]


def _rows(src: SheetSource, catalog: Catalog):
    summary = calculate_materials(src.cells, catalog, src.wastage_rate)
    beads = [
        [
            l.code,
            f"{l.name_zh} {l.name}",
            "MIYUKI Delica 11/0",
            l.required,
            l.reserve,
            l.purchase_needed,
            l.beads_per_package,
            l.packages if l.packages is not None else "需确认",
            "是" if src.bought.get(l.color_id) else "",
            "是" if src.out_of_stock.get(l.color_id) else "",
            "已核实" if l.verified else "示例数据",
        ]
        for l in summary.lines
    ]
    parts = [[a.name, a.quantity, a.reason, STATUS_ZH[a.status]] for a in estimate_accessories(src.item_type, src.width, src.height)]
    return summary, beads, parts


def shopping_csv(src: SheetSource, catalog: Catalog) -> bytes:
    _, beads, parts = _rows(src, catalog)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["Tsumii 材料清单", src.name])
    w.writerow(["珠型", "MIYUKI Delica 11/0", "损耗比例", f"{round(src.wastage_rate * 100)}%"])
    w.writerow([])
    w.writerow(BEAD_HEADER)
    w.writerows(beads)
    w.writerow([])
    w.writerow(PART_HEADER)
    w.writerows(parts)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def shopping_xlsx(src: SheetSource, catalog: Catalog) -> bytes:
    summary, beads, parts = _rows(src, catalog)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "珠子材料"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="ECEAE6")
    ws.append([f"Tsumii 材料清单 · {src.name}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["珠型", "MIYUKI Delica 11/0", "损耗比例", src.wastage_rate])
    ws["D2"].number_format = "0%"
    ws.append([])
    ws.append(BEAD_HEADER + ["色块"])
    for c in ws[4]:
        c.font, c.fill = bold, head_fill
    for i, (row, line) in enumerate(zip(beads, summary.lines), start=5):
        ws.append(row)
        # 采购量用公式表达，便于用户修改损耗比例后自动更新
        ws.cell(row=i, column=5, value=f"=CEILING(ROUND(D{i}*$D$2,9),1)")
        ws.cell(row=i, column=6, value=f"=D{i}+E{i}")
        if line.beads_per_package:
            ws.cell(row=i, column=8, value=f"=CEILING(F{i}/G{i},1)")
        ws.cell(row=i, column=12).fill = PatternFill("solid", fgColor=line.hex.lstrip("#"))
    last = 4 + len(beads)
    ws.append([])
    ws.append(["合计", "", "", f"=SUM(D5:D{last})", f"=SUM(E5:E{last})", f"=SUM(F5:F{last})", "", f"=SUM(H5:H{last})"])
    ws.cell(row=ws.max_row, column=1).font = bold
    for col, width in zip(range(1, 13), [14, 28, 18, 10, 10, 10, 10, 10, 8, 8, 10, 6]):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws2 = wb.create_sheet("其他配件")
    ws2.append(PART_HEADER)
    for c in ws2[1]:
        c.font, c.fill = bold, head_fill
    for r in parts:
        ws2.append(r)
    for col, width in zip("ABCD", [18, 12, 44, 10]):
        ws2.column_dimensions[col].width = width
        for cell in ws2[col]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
